"""Build the launch-plan outputs from content.py.

Usage: python3 tools/build_workbook.py
Outputs:
  deliverables/AI_Dad_Launch_Plan.xlsx   one tab per section
  deliverables/AI_Dad_Launch_Plan.csv    single tab, sections stacked (the Google Sheet layout)
  deliverables/sheet_paste.html          the same single-tab layout as an HTML table, for the in-place Sheets update
  docs/LAUNCH_PLAN.md                    readable copy
"""
import csv
import html
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(__file__))
import content  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "deliverables", "AI_Dad_Launch_Plan.xlsx")
CSV = os.path.join(ROOT, "deliverables", "AI_Dad_Launch_Plan.csv")
HTML = os.path.join(ROOT, "deliverables", "sheet_paste.html")
MD = os.path.join(ROOT, "docs", "LAUNCH_PLAN.md")

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def utf16_len(s):
    """Character count the way the apps' counters see it (an emoji counts as 2)."""
    return len(s.encode("utf-16-le")) // 2


def bio_copy(platform, field):
    for row in content.BIO[3]:
        if row[0] == platform and row[1] == field:
            return row[2]
    raise KeyError((platform, field))


SHORT = 90  # cells this short are already one fact; leave them whole
KEEP_WHOLE = {("Bio", 2)}  # bio copy must stay paste-exact
ABBREV = ("e.g.", "i.e.", "vs.", "a.m.", "p.m.", "No.", "Opp.", "etc.", "Cl.", "approx.", "St.", "U.S.", "Univ.", "Inc.")


def _split_top(text, is_cut):
    """Split text at positions where is_cut(text, i) is true, ignoring anything inside parentheses or quotes."""
    parts, depth, start = [], 0, 0
    for i, ch in enumerate(text):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        elif depth == 0 and is_cut(text, i):
            parts.append(text[start:i + 1].strip())
            start = i + 1
    parts.append(text[start:].strip())
    return [p for p in parts if p]


def _sentence_end(text, i):
    if text[i] not in ".!?" or i + 2 >= len(text) or text[i + 1] != " ":
        return False
    nxt = text[i + 2]
    if not (nxt.isupper() or nxt.isdigit() or nxt in "'(\"$#@~"):
        return False
    before = text[:i + 1]
    if any(before.endswith(a) for a in ABBREV):
        return False
    word = before.split()[-1] if before.split() else ""
    if word[:-1].isdigit() and before.strip() == word:  # "1. The ONE thing": a list number, not a sentence
        return False
    return not (len(word) <= 2 and word[0].isupper())  # initials like "J."


def split_facts(text):
    """One cell -> a list of short lines, one fact each. Text is never reworded, only split."""
    text = str(text)
    if len(text) <= SHORT and "\n" not in text:
        return [text]
    lines = []
    for chunk in text.split("\n"):
        chunk = chunk.strip()
        if not chunk:
            continue
        # Inline numbered lists: "(1) ... (2) ..."
        pieces = re.split(r"\s(?=\(\d+\)\s)", chunk)
        for piece in pieces:
            for sent in _split_top(piece, _sentence_end):
                # Arrow chains read better one step per line.
                if sent.count(" → ") >= 2 and len(sent) > SHORT:
                    steps = sent.split(" → ")
                    lines.append(steps[0])
                    lines.extend("→ " + st for st in steps[1:])
                    continue
                # Independent clauses joined by semicolons.
                semi = _split_top(sent, lambda t, i: t[i] == ";" and i + 1 < len(t) and t[i + 1] == " ")
                if len(sent) > SHORT and len(semi) > 1 and all(len(x) > 20 for x in semi):
                    parts = semi
                else:
                    parts = [sent]
                for part in parts:
                    lines.extend(_shorten(part))
    # Sheets swallows a leading apostrophe, so a line that opens with a quote gets a list dot in front.
    return ["· " + x if x.startswith("'") else x for x in lines] or [""]


LONG = 130  # a single line longer than this gets one more pass


def _shorten(line):
    """Second pass for a still-long line: '·' lists, a lead-in colon, then ', and/but/so' joins."""
    if len(line) <= LONG:
        return [line]
    if line.count(" · ") >= 2:
        head, *items = line.split(" · ")
        if ": " in head and len(head.split(": ", 1)[0]) < 60:
            lead, first = head.split(": ", 1)
            return [lead + ":"] + ["· " + x for x in [first] + items]
        return [head] + ["· " + x for x in items]
    depth = 0
    for i, ch in enumerate(line):
        depth += ch == "("
        depth -= ch == ")" and depth > 0
        if depth == 0 and line.startswith(": ", i) and 12 <= i <= 130 and "http" not in line[max(0, i - 6):i]:
            return [line[:i + 1]] + _shorten(line[i + 2:])
    cut = None
    depth = 0
    for i, ch in enumerate(line):
        depth += ch == "("
        depth -= ch == ")" and depth > 0
        if depth == 0 and i > 40 and len(line) - i > 30 and any(line.startswith(j, i) for j in (", and ", ", but ", ", so ", ", because ", ", which ")):
            if cut is None or abs(i - len(line) / 2) < abs(cut - len(line) / 2):
                cut = i
    if cut is not None:
        return _shorten(line[:cut + 1]) + _shorten(line[cut + 2:])
    # Last resort for a long comma list: one item per line.
    items = _split_top(line, lambda t, i: t[i] == "," and t.startswith(", ", i))
    # A quote mark not between two letters (so not "I'll") opens or closes a quote; never cut inside one.
    quotes = lambda x: len(re.findall(r"(?<![A-Za-z])'|'(?![A-Za-z])", x))
    if len(items) >= 3 and all(len(x) > 12 and quotes(x) % 2 == 0 for x in items):
        return [items[0]] + ["· " + x for x in items[1:]]
    return [line]


def stacked_rows(tabs):
    """The single-tab Google Sheet layout: every section stacked, one short fact per row.

    A record whose long cells split into several lines spreads down over several rows (each column reads
    straight down), followed by a blank spacer row so records stay visually separate."""
    rows = [["AI DAD HEALTHY DAD: LAUNCH PLAN"],
            ["Sections, in order:"]]
    rows.extend(["", "· " + t[0]] for t in tabs)
    for title, intro, header, body in tabs:
        rows.append([])
        rows.append([f"■ {title.upper()}"])
        for line in intro:
            rows.extend([f] for f in split_facts(line))
        rows.append(list(header))
        for record in body:
            cols = [[str(c)] if (title, j) in KEEP_WHOLE else split_facts(c) for j, c in enumerate(record)]
            height = max(len(c) for c in cols)
            for k in range(height):
                rows.append([c[k] if k < len(c) else "" for c in cols])
            if height > 1:
                rows.append([])
    while rows and not rows[-1]:
        rows.pop()
    return rows


def section_layout(rows):
    """Row indexes (0-based) of section-title rows and table-header rows, for sheet formatting."""
    titles = [i for i, r in enumerate(rows) if r and str(r[0]).startswith("■ ")]
    headers = []
    for t in titles:
        i = t + 1
        while i < len(rows) and len(rows[i]) < 2:
            i += 1
        headers.append(i)
    return titles, headers


def fill_placeholders(tabs):
    """Replace {ROWS} and the bio character counts with real values."""
    values = {
        "{TT_BIO_LEN}": str(utf16_len(bio_copy("TikTok", "Bio"))),
        "{IG_NAME_LEN}": str(utf16_len(bio_copy("Instagram", "Name"))),
        "{IG_BIO_LEN}": str(utf16_len(bio_copy("Instagram", "Bio"))),
        "{TT_TNAME_LEN}": str(utf16_len(bio_copy("TikTok", "Name (teaser phase, to Nov 3)"))),
        "{TT_TBIO_LEN}": str(utf16_len(bio_copy("TikTok", "Bio (teaser phase, to Nov 3)"))),
        "{IG_TNAME_LEN}": str(utf16_len(bio_copy("Instagram", "Name (teaser phase, to Nov 3)"))),
        "{IG_TBIO_LEN}": str(utf16_len(bio_copy("Instagram", "Bio (teaser phase, to Nov 3)"))),
        "{ROWS}": "000",  # set below, once the layout is known
    }

    def fill(cell):
        cell = str(cell)
        for key, val in values.items():
            cell = cell.replace(key, val)
        return cell

    # A 3-digit row count never changes how a cell splits, so one dry run gives the final count.
    dry = [(fill(t), [fill(i) for i in intro], [fill(h) for h in header], [[fill(c) for c in r] for r in body])
           for t, intro, header, body in tabs]
    values["{ROWS}"] = str(len(stacked_rows(dry)))
    return [(fill(t), [fill(i) for i in intro], [fill(h) for h in header], [[fill(c) for c in r] for r in body])
            for t, intro, header, body in tabs], values


def col_width(header, rows, i):
    longest = max([len(str(header[i]))] + [min(len(str(r[i])), 400) for r in rows])
    return max(10, min(70, int(longest / 3.2) + 8)) if longest > 30 else max(10, longest + 4)


def build_xlsx(tabs):
    wb = Workbook()
    wb.remove(wb.active)
    for title, intro, header, rows in tabs:
        ws = wb.create_sheet(title[:31])
        r = 1
        for line in intro:
            ws.cell(row=r, column=1, value=line).font = Font(italic=True)
            r += 1
        r += 1 if intro else 0
        header_row = r
        for c, h in enumerate(header, 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
        for row in rows:
            r += 1
            for c, v in enumerate(row, 1):
                # Stored as string cells so Sheets keeps "1", "#" etc. as text on import.
                cell = ws.cell(row=r, column=c, value=str(v))
                cell.alignment = WRAP
                cell.data_type = "s"
        for i in range(len(header)):
            ws.column_dimensions[chr(65 + i)].width = col_width(header, rows, i)
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)


def build_csv(rows):
    with open(CSV, "w", newline="") as f:
        csv.writer(f).writerows(rows)


def build_html(rows):
    """One-line HTML table. No double quotes, backslashes or raw newlines survive,
    so the string can sit inside a JSON body untouched. Rows aren't padded: the
    update clears the sheet first, so short rows can't leave stale cells behind."""

    def td(v):
        # Only &, <, > and " need escaping; apostrophes are safe in both HTML text and JSON.
        v = html.escape(str(v), quote=False).replace('"', "&quot;").replace("\n", "<br>")
        return f"<td>{v}</td>"

    out = ["<table>"]
    for r in rows:
        cells = list(r) or [""]
        out.append("<tr>" + "".join(td(c) for c in cells) + "</tr>")
    out.append("</table>")
    page = "".join(out)
    assert '"' not in page and "\\" not in page and "\n" not in page
    with open(HTML, "w") as f:
        f.write(page)
    return page


def md_cell(v):
    return str(v).replace("\n", "<br>")


def build_md(tabs):
    out = ["# AI Dad Healthy Dad: Launch Plan", "",
           "Generated from `tools/content.py`. The same content lives in the project Google Sheet.", ""]
    for title, intro, header, rows in tabs:
        out += [f"## {title}", ""]
        out += [f"> {line}" for line in intro]
        if intro:
            out.append("")
        for row in rows:
            first = md_cell(row[0])
            if title == "Sources":
                out.append(f"- **{first}**: {row[1]}: {row[2]}")
                continue
            out.append(f"### {first}")
            for h, v in zip(header[1:], row[1:]):
                if str(v).strip() and str(v) != "-":
                    out.append(f"- **{h}:** {md_cell(v)}")
            out.append("")
        out.append("")
    os.makedirs(os.path.dirname(MD), exist_ok=True)
    with open(MD, "w") as f:
        f.write("\n".join(out))


def build():
    tabs, values = fill_placeholders(content.TABS)
    rows = stacked_rows(tabs)
    build_xlsx(tabs)
    build_csv(rows)
    build_html(rows)
    build_md(tabs)
    return tabs, rows, values


if __name__ == "__main__":
    _, rows, values = build()
    print("rows in sheet layout:", len(rows), "| placeholders:", values)
    for path in (XLSX, CSV, HTML, MD):
        print(path, os.path.getsize(path), "bytes")
